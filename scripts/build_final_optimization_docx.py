#!/usr/bin/env python3
"""Wrap verified final report pages in a fixed-layout Word file."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.shared import Inches, Pt

ROOT = Path(__file__).resolve().parents[1]
PAGE_DIR = ROOT / "reports" / "final_optimization_pages_v2"
OUTPUT = ROOT / "reports" / "CTCPGRN_final_optimized_results.docx"
pages = sorted(PAGE_DIR.glob("page-*.png"), key=lambda p: int(p.stem.split("-")[-1]))
if not pages:
    raise SystemExit("No report pages found")

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.5)
section.page_height = Inches(11)
section.top_margin = Inches(0.22)
section.bottom_margin = Inches(0.22)
section.left_margin = Inches(0.25)
section.right_margin = Inches(0.25)
section.header_distance = Inches(0)
section.footer_distance = Inches(0)
for index, page in enumerate(pages):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(page), width=Inches(8.0))
    if index < len(pages)-1:
        paragraph.add_run().add_break(WD_BREAK.PAGE)
doc.core_properties.title = "连续时间变化点基因调控网络模型分优先级优化与最终独立验证报告"
doc.core_properties.subject = "CT-CP-GRN final optimisation and validation"
doc.core_properties.author = "CT-CP-GRN research project"
doc.save(OUTPUT)
print(OUTPUT)
