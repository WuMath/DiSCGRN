from pathlib import Path

from docx import Document
from docx.oxml.ns import qn

doc = Document()
for font in [
    "Arial Unicode MS", "Songti SC", "PingFang SC", "Heiti SC", "STHeiti",
    "Hiragino Sans GB", "Microsoft YaHei", "SimSun", "Apple LiGothic",
]:
    p = doc.add_paragraph()
    run = p.add_run(f"{font}: 连续时间变化点基因调控网络 参数校准 稳定性评估")
    run.font.name = font
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
doc.save(Path("/tmp/font_test.docx"))
