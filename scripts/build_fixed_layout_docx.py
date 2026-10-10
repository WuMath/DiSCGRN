#!/usr/bin/env python3
"""Wrap verified report pages in a fixed-layout Word document."""

from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK, WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Mm, Pt


ROOT = Path(__file__).resolve().parents[1]
PAGE_DIR = ROOT / "reports" / "fixed_final"
OUTPUT = ROOT / "reports" / "CTCPGRN_calibrated_case_study.docx"

pages = sorted(PAGE_DIR.glob("page-*.png"), key=lambda path: int(path.stem.split("-")[-1]))
if not pages:
    raise SystemExit("No rendered pages found")

doc = Document()
section = doc.sections[0]
section.page_width = Mm(210)
section.page_height = Mm(297)
section.top_margin = Mm(8)
section.bottom_margin = Mm(8)
section.left_margin = Mm(10)
section.right_margin = Mm(10)
section.header_distance = Mm(0)
section.footer_distance = Mm(0)

for index, page in enumerate(pages):
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.add_run().add_picture(str(page), width=Mm(190))
    if index < len(pages) - 1:
        paragraph.add_run().add_break(WD_BREAK.PAGE)

doc.core_properties.title = "连续时间变化点GRN推断：参数校准、基线比较与稳定性评估"
doc.core_properties.subject = "DREAM4非平稳完整算例（固定版式）"
doc.core_properties.author = "CT–CP–GRN研究项目"
doc.save(OUTPUT)
print(OUTPUT)
