#!/usr/bin/env python3
"""Build the calibrated DREAM4 case-study report as a publication-style DOCX."""

from __future__ import annotations

import argparse
import json
from pathlib import Path

import numpy as np
import pandas as pd

try:
    from docx import Document
    from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Cm, Inches, Pt, RGBColor
except ModuleNotFoundError:
    # Figure-only mode runs in the scientific environment; document mode uses
    # the bundled artifact environment where python-docx is available.
    pass


ROOT = Path(__file__).resolve().parents[1]
RESULTS = ROOT / "results" / "calibrated_case_study"
FIGURES = ROOT / "reports" / "figures"
OUTPUT = ROOT / "reports" / "CTCPGRN_calibrated_case_study.docx"
FIGURES.mkdir(parents=True, exist_ok=True)
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

COLORS = {"ct": "#214E76", "euler": "#D17A22", "dplsq": "#3B7A57", "two_stage": "#8A5A9B"}
LABELS = {"ct": "CT–CP–GRN", "euler": "Euler-DP", "dplsq": "DPLSQ-TV", "two_stage": "Two-stage"}


def configure_plots() -> None:
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "font.family": "sans-serif",
        "font.sans-serif": ["Arial Unicode MS", "PingFang SC", "DejaVu Sans"],
        "axes.unicode_minus": False,
        "font.size": 9,
        "figure.dpi": 180,
    })


def save_figures(calibration: pd.DataFrame, test: pd.DataFrame, bootstrap: pd.DataFrame, summary: dict) -> dict[str, Path]:
    import matplotlib.pyplot as plt
    configure_plots()
    paths = {}

    fig, axes = plt.subplots(2, 2, figsize=(9.2, 6.6), constrained_layout=True)
    alphas = sorted(calibration.alpha.unique())
    betas = sorted(calibration.beta.unique())
    for ax, method in zip(axes.flat, ["ct", "euler", "dplsq", "two_stage"]):
        sub = calibration[calibration.method == method].pivot(index="alpha", columns="beta", values="selection_score").reindex(index=alphas, columns=betas)
        image = ax.imshow(sub.values, cmap="YlGnBu", aspect="auto", vmin=0, vmax=max(1.12, calibration.selection_score.max()))
        ax.set_title(LABELS[method], fontweight="bold")
        ax.set_xticks(range(len(betas)), [f"{x:g}" for x in betas])
        ax.set_yticks(range(len(alphas)), [f"{x:g}" for x in alphas])
        ax.set_xlabel("变化点惩罚 β")
        ax.set_ylabel("稀疏惩罚 α")
        for i in range(len(alphas)):
            for j in range(len(betas)):
                ax.text(j, i, f"{sub.values[i, j]:.2f}", ha="center", va="center", color="white" if sub.values[i, j] > .62 else "#222")
    fig.colorbar(image, ax=axes.ravel().tolist(), shrink=.78, label="校准选择得分")
    paths["grid"] = FIGURES / "calibration_grid.png"
    fig.savefig(paths["grid"], bbox_inches="tight")
    plt.close(fig)

    order = ["ct", "dplsq", "euler", "two_stage"]
    test = test.set_index("method").loc[order].reset_index()
    fig, axes = plt.subplots(1, 2, figsize=(9.2, 3.5), constrained_layout=True)
    x = np.arange(len(order))
    axes[0].bar(x, test.estimated_change_time, color=[COLORS[m] for m in order])
    axes[0].axhline(float(summary["true_change_time"]), color="#B22222", ls="--", lw=1.6, label="真实变化时刻")
    axes[0].set_xticks(x, [LABELS[m] for m in order], rotation=18, ha="right")
    axes[0].set_ylabel("变化时刻")
    axes[0].set_title("独立测试：变化点定位", fontweight="bold")
    axes[0].legend(frameon=False)
    width = .36
    axes[1].bar(x - width/2, test.network_aupr, width, label="网络AUPR", color="#497AA1")
    axes[1].bar(x + width/2, test.change_aupr, width, label="变化边AUPR", color="#A36B3F")
    axes[1].set_xticks(x, [LABELS[m] for m in order], rotation=18, ha="right")
    axes[1].set_ylim(0, .36)
    axes[1].set_ylabel("AUPR")
    axes[1].set_title("独立测试：边级恢复", fontweight="bold")
    axes[1].legend(frameon=False)
    paths["test"] = FIGURES / "independent_test.png"
    fig.savefig(paths["test"], bbox_inches="tight")
    plt.close(fig)

    counts = bootstrap.change_time.value_counts().sort_index()
    fig, ax = plt.subplots(figsize=(8.7, 3.6), constrained_layout=True)
    bars = ax.bar([f"{x:.3f}" for x in counts.index], counts.values, color="#214E76")
    true_time = float(summary["true_change_time"])
    true_label = f"{true_time:.3f}"
    for bar, label in zip(bars, [f"{x:.3f}" for x in counts.index]):
        if label == true_label:
            bar.set_color("#B22222")
        ax.text(bar.get_x() + bar.get_width()/2, bar.get_height() + 2, str(int(bar.get_height())), ha="center", fontsize=8)
    ax.set_xlabel("bootstrap估计变化时刻")
    ax.set_ylabel("频数（共200次）")
    ax.set_title("CT–CP–GRN变化点稳定性", fontweight="bold")
    ax.spines[["top", "right"]].set_visible(False)
    paths["bootstrap"] = FIGURES / "bootstrap_distribution.png"
    fig.savefig(paths["bootstrap"], bbox_inches="tight")
    plt.close(fig)
    return paths


def shade_cell(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def add_table(document: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    set_repeat_table_header(table.rows[0])
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        cell.text = header
        shade_cell(cell, "214E76")
        for run in cell.paragraphs[0].runs:
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.bold = True
            run.font.size = Pt(8.5)
        if widths:
            cell.width = Inches(widths[i])
    for row_values in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row_values):
            cells[i].text = value
            cells[i].vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if widths:
                cells[i].width = Inches(widths[i])
            for paragraph in cells[i].paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    run.font.size = Pt(8.2)
    document.add_paragraph()


def add_caption(document: Document, text: str) -> None:
    p = document.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.style = document.styles["Caption"]


def set_cell_margins(table) -> None:
    for row in table.rows:
        for cell in row.cells:
            tc = cell._tc
            tc_pr = tc.get_or_add_tcPr()
            margins = tc_pr.first_child_found_in("w:tcMar")
            if margins is None:
                margins = OxmlElement("w:tcMar")
                tc_pr.append(margins)
            for edge in ("top", "left", "bottom", "right"):
                element = margins.find(qn(f"w:{edge}"))
                if element is None:
                    element = OxmlElement(f"w:{edge}")
                    margins.append(element)
                element.set(qn("w:w"), "65")
                element.set(qn("w:type"), "dxa")


def add_equation(document: Document, equation: str) -> None:
    p = document.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(equation)
    run.font.name = "Cambria Math"
    run.font.size = Pt(10.5)


def force_cjk_font(document: Document) -> None:
    """Apply an explicit CJK font and language to every run containing Han text."""
    paragraphs = list(document.paragraphs)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                paragraphs.extend(cell.paragraphs)
    for section in document.sections:
        paragraphs.extend(section.header.paragraphs)
        paragraphs.extend(section.footer.paragraphs)
    for paragraph in paragraphs:
        for run in paragraph.runs:
            if not any("\u3400" <= char <= "\u9fff" for char in run.text):
                continue
            run.font.name = "Arial Unicode MS"
            r_pr = run._element.get_or_add_rPr()
            r_fonts = r_pr.rFonts
            if r_fonts is None:
                r_fonts = OxmlElement("w:rFonts")
                r_pr.insert(0, r_fonts)
            for key in ("ascii", "hAnsi", "eastAsia", "cs"):
                r_fonts.set(qn(f"w:{key}"), "Arial Unicode MS")
            language = r_pr.find(qn("w:lang"))
            if language is None:
                language = OxmlElement("w:lang")
                r_pr.append(language)
            language.set(qn("w:val"), "zh-CN")
            language.set(qn("w:eastAsia"), "zh-CN")


def build_document(calibration: pd.DataFrame, selected: pd.DataFrame, test: pd.DataFrame, bootstrap: pd.DataFrame, summary: dict, figures: dict[str, Path]) -> None:
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Cm(2.1)
    section.bottom_margin = Cm(2.0)
    section.left_margin = Cm(2.25)
    section.right_margin = Cm(2.25)
    section.header_distance = Cm(1.0)
    section.footer_distance = Cm(1.0)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial Unicode MS"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.35
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.first_line_indent = Pt(21)
    for name, size, color in [("Title", 21, "17365D"), ("Heading 1", 15, "17365D"), ("Heading 2", 12, "365F91"), ("Heading 3", 10.5, "365F91")]:
        style = styles[name]
        style.font.name = "Arial Unicode MS"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)
    styles["Caption"].font.name = "Arial Unicode MS"
    styles["Caption"]._element.rPr.rFonts.set(qn("w:eastAsia"), "Arial Unicode MS")
    styles["Caption"].font.size = Pt(8.5)
    styles["Caption"].font.italic = False
    styles["Caption"].font.color.rgb = RGBColor.from_string("444444")

    header = section.header.paragraphs[0]
    header.text = "CT–CP–GRN｜参数校准与完整算例报告"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.runs[0].font.size = Pt(8)
    header.runs[0].font.color.rgb = RGBColor.from_string("777777")
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("第 ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    footer.add_run(" 页")

    title = doc.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("连续时间变化点GRN推断：\n参数校准、基线比较与稳定性评估")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.first_line_indent = Pt(0)
    run = subtitle.add_run("DREAM4 size-10 network 1 非平稳完整算例｜内部实验报告")
    run.bold = True
    run.font.color.rgb = RGBColor.from_string("365F91")
    doc.add_paragraph()
    meta = doc.add_table(rows=4, cols=2)
    meta.style = "Light Shading Accent 1"
    meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    for row, (key, value) in zip(meta.rows, [
        ("分析日期", "2026-09-29"),
        ("独立测试种子", "91001（未参与参数选择）"),
        ("校准规模", "4种方法 × 3个α × 4个β × 3个独立种子 = 144次拟合组合"),
        ("不确定性分析", "CT–CP–GRN按重复实验重采样，200次bootstrap"),
    ]):
        row.cells[0].text, row.cells[1].text = key, value
    doc.add_paragraph()
    callout = doc.add_table(rows=1, cols=1)
    callout.style = "Light Shading Accent 1"
    cell = callout.cell(0, 0)
    shade_cell(cell, "EAF2F8")
    cell.text = (
        "核心结论｜在完全独立的测试模拟中，CT–CP–GRN与DPLSQ-TV均精确定位真实变化时刻；"
        "CT–CP–GRN的200次bootstrap有43.5%精确命中，77.0%落在真实点或其前一相邻采样点。"
        "但全部方法的边级AUPR仅为0.174–0.300，因此本结果支持“变化点定位可行”，尚不足以宣称高精度GRN边恢复。"
    )
    cell.paragraphs[0].paragraph_format.first_line_indent = Pt(0)
    cell.paragraphs[0].runs[0].font.bold = True
    doc.add_page_break()

    doc.add_heading("摘要", level=1)
    doc.add_paragraph(
        "本报告完成一个可重复的非平稳DREAM4完整算例：先用三个独立随机种子校准稀疏惩罚α与变化点惩罚β，"
        "再在未参与校准的测试种子上比较CT–CP–GRN、Euler-DP、DPLSQ-TV与Two-stage；随后实施严格的时间阻断交叉验证，"
        "并对主模型进行200次重复实验层面的bootstrap。连续时间主模型使用矩阵指数精确传播，不把不规则时间间隔近似为固定步长；"
        "动态规划在候选观测时刻间选择零个或一个变化点。独立测试中，CT–CP–GRN与DPLSQ-TV准确命中真实时刻6.303，"
        "Euler-DP和Two-stage分别误判为9.628和2.708。DPLSQ-TV的阻断预测RMSE最低（0.0373），而CT–CP–GRN的网络AUPR"
        "高于DPLSQ-TV和Euler-DP，但仍只有0.2203。bootstrap显示估计主要集中在5.483与6.303两个相邻时间点。"
        "这些结果构成方法实现与评估流程的端到端验证，但单网络、单噪声场景不构成性能优越性的充分证据。"
    )
    doc.add_heading("1　研究问题与本轮实验边界", level=1)
    doc.add_paragraph(
        "目标不是用一个有利案例替代完整论文实验，而是回答三个更基础的问题：连续时间分段模型能否从不规则采样表达轨迹中恢复结构变化时刻；"
        "相同搜索框架下，连续时间传播相对于Euler近似和导数回归有什么表现；变化点结论对生物重复抽样是否稳定。"
        "因此本轮只使用DREAM4 size-10 network 1的已知有向符号网络生成一个非平稳情景，固定测试种子后不再据测试结果调参。"
    )
    add_table(doc, ["项目", "预注册设置", "意义"], [
        ["网络", "DREAM4 size-10 network 1", "保留已知有向符号拓扑作为边级真值"],
        ["非平稳机制", "35%边重连；10%边符号翻转；权重扰动尺度0.35", "产生可解释的阶段特异性网络"],
        ["采样", "13个不规则时点，时间范围0–12；5个重复", "检验对非均匀Δt的适应性"],
        ["真实变化点", "索引6，t*=6.303", "位于中段，保证两侧均有足够转移"],
        ["观测噪声", "高斯噪声标准差0.03", "形成有限样本识别难度"],
        ["候选约束", "每段至少3个状态转移", "避免极短分段的过拟合"],
    ], [1.25, 2.35, 2.8])

    doc.add_heading("2　模型、变量与目标函数", level=1)
    doc.add_heading("2.1 连续时间分段GRN", level=2)
    doc.add_paragraph(
        "令x(t)∈ℝᵖ表示p个基因在连续时间t的表达状态。变化点集合把时间轴划分为若干区间；在区间k内，"
        "动力学矩阵Aₖ保持不变，Aₖ,ij表示基因j对基因i的瞬时线性调控效应，正值与负值分别代表激活样与抑制样效应；"
        "bₖ是区间特异的基线输入。对本轮单变化点问题，k∈{1,2}。"
    )
    add_equation(doc, "dx(t)/dt = Aₖx(t) + bₖ,    τₖ₋₁ ≤ t < τₖ")
    doc.add_paragraph(
        "对相邻观测(tᵢ,xᵢ)与(tᵢ₊₁,xᵢ₊₁)，令Δtᵢ=tᵢ₊₁−tᵢ。CT–CP–GRN不采用Euler离散，"
        "而使用线性常微分方程的精确转移；积分项同时传播常数基线。这样每个不规则时间间隔都直接进入转移算子。"
    )
    add_equation(doc, "x̂ᵢ₊₁ = exp(AₖΔtᵢ)xᵢ + ∫₀^Δtᵢ exp(Aₖs)bₖ ds")
    doc.add_heading("2.2 分段代价与动态规划", level=2)
    doc.add_paragraph(
        "给定区间I，数据项衡量精确ODE预测与观测的加权平方误差；α控制非对角调控边的稀疏度。"
        "对角元素描述自衰减或自反馈，不纳入边稀疏惩罚。为保证不同长度分段可相加，平均数据损失按区间转移数nᴵ还原为总损失，"
        "而稀疏惩罚在每个区间只计一次。"
    )
    add_equation(doc, "C(I;α) = Σᵢ∈I ||Wᵢ¹ᐟ²(xᵢ₊₁−x̂ᵢ₊₁)||₂² + α||Aᴵ_off||₁")
    doc.add_paragraph(
        "外层目标在无变化模型与所有合法单变化候选之间取最小值。β是新增一个变化点的固定复杂度代价；"
        "β越大，模型越倾向保持平稳。此处候选仅在实际采样时刻上取值，因此搜索是有限离散优化。"
    )
    add_equation(doc, "min { C([1,T];α),  min₍τ₎ [C([1,τ];α)+C([τ,T];α)+β] }")
    doc.add_heading("2.3 三个对照方法", level=2)
    add_table(doc, ["方法", "状态转移/变化点策略", "对照目的"], [
        ["Euler-DP", "xᵢ₊₁≈xᵢ+Δtᵢ(Axᵢ+b)；与主模型相同的候选扫描", "隔离精确连续传播与一阶离散近似的差异"],
        ["DPLSQ-TV", "用(xᵢ₊₁−xᵢ)/Δtᵢ回归xᵢ；同样执行离散分段搜索", "检验数值求导噪声与快速稀疏回归基线"],
        ["Two-stage", "先按表达均值跃迁的SSE选变化点，再分别拟合连续时间网络", "检验变化点与网络联合估计是否必要"],
    ], [1.15, 3.2, 2.15])

    doc.add_heading("3　参数校准与防止信息泄漏", level=1)
    doc.add_paragraph(
        "参数校准仅使用种子4101、4102与4103生成的三套独立模拟；测试种子91001在参数冻结后才运行。"
        "完整网格包含α∈{0.0005,0.002,0.008}与β∈{0,0.002,0.01,0.04}，每种方法12组、四种方法共48组参数单元。"
        "每个单元跨三种子取均值，并用下式排序；第一项优先保证变化点检出，后两项仅在变化点表现接近时区分边恢复。"
    )
    add_equation(doc, "S = mean(F1_CP) + 0.25·mean(AUPR_change) + 0.10·mean(AUPR_network)")
    doc.add_paragraph(
        "若选择得分完全相同，先取较小β，再取较大α。该规则在查看测试结果前固定。β=0被选中说明本模拟中分段拟合收益已足够明显；"
        "这不意味着真实数据也应取消变化点惩罚，真实分析仍需用嵌套交叉验证或信息准则校准β。"
    )
    doc.add_picture(str(figures["grid"]), width=Inches(6.55))
    add_caption(doc, "图1　完整参数网格的校准选择得分。每格为3个独立校准种子的平均综合得分。")
    selected_rows = []
    for row in selected.itertuples(index=False):
        selected_rows.append([
            LABELS[row.method], f"{row.alpha:g}", f"{row.beta:g}", f"{row.mean_cp_f1:.3f}",
            f"{row.mean_change_aupr:.3f}", f"{row.mean_network_aupr:.3f}", f"{row.selection_score:.3f}",
        ])
    add_table(doc, ["方法", "α", "β", "CP F1", "变化边AUPR", "网络AUPR", "选择得分"], selected_rows, [.95, .52, .52, .62, 1.0, .83, .82])

    doc.add_heading("4　评估设计", level=1)
    doc.add_heading("4.1 独立测试指标", level=2)
    doc.add_paragraph(
        "变化点在±1.5时间单位容差内视为命中，报告precision、recall、F1及命中后的绝对时间误差。"
        "网络恢复分别计算两个阶段有向边分数相对于真值的AUROC与AUPR，再取阶段平均；变化边指标以|Â₂−Â₁|"
        "排序，对照真实发生改变的边。由于正边比例低，AUPR比AUROC更有解释力。"
    )
    doc.add_heading("4.2 时间阻断交叉验证", level=2)
    doc.add_paragraph(
        "为避免相邻时间点泄漏，本轮把状态转移而非单个观测随机拆分。依次留出转移块{1,2}、{5,6}、{9,10}，"
        "每个训练转移作为独立两点轨迹进入拟合，防止删除区间后人为跨越时间缺口。变化点位置固定为该方法在完整测试轨迹上的估计，"
        "交叉验证只衡量条件于该分段的外推误差，因此不能替代嵌套的变化点选择验证。"
    )
    doc.add_heading("4.3 200次bootstrap", level=2)
    doc.add_paragraph(
        "每次从5个生物学重复中有放回抽取5个完整重复，重新计算时点均值与方差权重，并用已冻结的CT–CP–GRN参数重新搜索变化点。"
        "该做法保持每个重复内部的时间相关结构；若逐时点独立重采样，会破坏轨迹配对并夸大有效样本量。"
    )

    doc.add_heading("5　结果", level=1)
    test_rows = []
    for method in ["ct", "dplsq", "euler", "two_stage"]:
        row = test[test.method == method].iloc[0]
        err = abs(row.estimated_change_time - row.true_change_time)
        test_rows.append([
            LABELS[method], f"{row.estimated_change_time:.3f}", f"{err:.3f}", f"{row.cp_f1:.2f}",
            f"{row.blocked_cv_rmse_mean:.4f}±{row.blocked_cv_rmse_sd:.4f}", f"{row.network_aupr:.3f}", f"{row.change_aupr:.3f}",
        ])
    add_table(doc, ["方法", "估计τ", "|误差|", "CP F1", "blocked CV RMSE", "网络AUPR", "变化边AUPR"], test_rows, [1.0, .7, .65, .62, 1.35, .86, .96])
    doc.add_paragraph(
        "CT–CP–GRN和DPLSQ-TV精确命中真实变化时刻6.303。Euler-DP选择9.628，表明一阶传播误差在不规则、"
        "跨度较大的时间间隔上可能改变分段代价的排序；Two-stage选择2.708，说明仅凭表达均值跃迁寻找边界会把早期动力学松弛误当作网络重塑。"
        "DPLSQ-TV得到最低预测RMSE，但其网络AUPR低于CT–CP–GRN；这体现预测准确与结构识别不是同一目标。"
    )
    doc.add_picture(str(figures["test"]), width=Inches(6.55))
    add_caption(doc, "图2　独立测试集的变化点定位与边级恢复。红色虚线为真实变化时刻。")

    exact = int((np.isclose(bootstrap.change_time, float(summary["true_change_time"]))).sum())
    previous_time = float(summary["times"][summary["true_change_index"] - 1])
    adjacent = int((np.isclose(bootstrap.change_time, previous_time)).sum())
    within = int((np.abs(bootstrap.change_time - float(summary["true_change_time"])) <= 1.5).sum())
    doc.add_heading("5.2 Bootstrap稳定性", level=2)
    doc.add_paragraph(
        f"200次重采样中，{exact}次（{exact/2:.1f}%）精确命中6.303，{adjacent}次（{adjacent/2:.1f}%）落在前一相邻时刻5.483；"
        f"共有{within}次（{within/2:.1f}%）位于预设±1.5容差内。其余估计主要分布在2.708与7.156，"
        "提示五个重复仍不足以把变化点不确定性压缩到单一采样时刻。最稳妥的报告方式是给出离散频率分布或相邻时点区间，而非只报告点估计。"
    )
    doc.add_picture(str(figures["bootstrap"]), width=Inches(6.35))
    add_caption(doc, "图3　CT–CP–GRN的200次重复层bootstrap变化点分布。红柱为真实变化时刻。")

    doc.add_heading("6　结论、局限与下一步", level=1)
    doc.add_heading("6.1 可以支持的结论", level=2)
    for text in [
        "实现层面：精确连续时间传播、离散变化点搜索、三类基线、时间阻断验证与重复层bootstrap已经形成闭环，能够复现实验。",
        "算例层面：主模型在独立测试中精确定位变化点，且bootstrap的大部分质量集中在真实点及其相邻前一时点。",
        "方法学层面：连续时间建模在本例中改善了相对于Euler-DP的变化点定位，但DPLSQ-TV同样命中且预测误差更低，不能据此宣称主模型全面优越。",
    ]:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Pt(0)
        p.add_run(text)
    doc.add_heading("6.2 不能支持的结论", level=2)
    doc.add_paragraph(
        "本实验只有一个10基因网络、一个非平稳机制和一个测试种子；边级AUPR接近类别基率，且主模型校准阶段仅2/3种子命中。"
        "因此不能把本报告作为算法显著优于基线、可扩展至百基因、或能在真实数据中恢复因果调控的证据。"
    )
    doc.add_heading("6.3 投稿前必须补齐", level=2)
    add_table(doc, ["优先级", "补充实验", "最低建议"], [
        ["P0", "多网络、多机制模拟", "5个DREAM4 size-10网络；≥20种子；改变噪声、重复数、采样不规则度、重连比例"],
        ["P0", "严格嵌套模型选择", "外层按连续时间块测试；内层选择α、β与变化点数；禁止测试集参与阈值选择"],
        ["P1", "零变化与多变化场景", "同时报告假阳性率、变化点数量误差、Hausdorff距离"],
        ["P1", "真实数据验证", "GSE74243做阶段发现；CollecTRI/ChIP-Atlas仅作外部边证据，不当作完整真值"],
        ["P1", "统计比较", "跨网络/种子的配对bootstrap或Wilcoxon检验，并报告效应量与置信区间"],
    ], [.65, 2.0, 4.0])

    doc.add_heading("7　可重复性清单", level=1)
    doc.add_paragraph("本轮所有中间结果均保留为纯文本表格，Word只读取这些冻结结果，不在排版阶段重新拟合。")
    files = [
        "results/calibrated_case_study/calibration_grid.tsv（完整48格校准结果）",
        "results/calibrated_case_study/selected_parameters.tsv（冻结参数）",
        "results/calibrated_case_study/test_results.tsv（独立测试与blocked CV）",
        "results/calibrated_case_study/bootstrap_200.tsv（200次bootstrap逐次结果）",
        "results/calibrated_case_study/summary.json（模拟设置与时间网格）",
        "scripts/run_calibrated_case_study.py（运行入口）",
        "src/ctcpgrn/baselines.py（基线与变化点扫描实现）",
    ]
    for item in files:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.first_line_indent = Pt(0)
        p.add_run(item)
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Pt(0)
    run = p.add_run("复现命令：")
    run.bold = True
    code = p.add_run(" .venv/bin/python scripts/run_calibrated_case_study.py --bootstrap 200 --max-iter 25")
    code.font.name = "Menlo"
    code.font.size = Pt(8.5)

    for table in doc.tables:
        set_cell_margins(table)
    doc.core_properties.title = "连续时间变化点GRN推断：参数校准、基线比较与稳定性评估"
    doc.core_properties.subject = "DREAM4非平稳完整算例"
    doc.core_properties.author = "CT–CP–GRN研究项目"
    force_cjk_font(doc)
    doc.save(OUTPUT)


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--mode", choices=["all", "figures", "document"], default="all")
    args = parser.parse_args()
    calibration = pd.read_csv(RESULTS / "calibration_grid.tsv", sep="\t")
    selected = pd.read_csv(RESULTS / "selected_parameters.tsv", sep="\t")
    test = pd.read_csv(RESULTS / "test_results.tsv", sep="\t")
    bootstrap = pd.read_csv(RESULTS / "bootstrap_200.tsv", sep="\t")
    summary = json.loads((RESULTS / "summary.json").read_text())
    figures = {
        "grid": FIGURES / "calibration_grid.png",
        "test": FIGURES / "independent_test.png",
        "bootstrap": FIGURES / "bootstrap_distribution.png",
    }
    if args.mode in {"all", "figures"}:
        figures = save_figures(calibration, test, bootstrap, summary)
    if args.mode in {"all", "document"}:
        build_document(calibration, selected, test, bootstrap, summary, figures)
        print(OUTPUT)


if __name__ == "__main__":
    main()
